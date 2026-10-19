import streamlit as st
import pandas as pd
import pypdf
import re
from datetime import datetime
import io
import csv
import requests
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Liquors Ventures Analytics", layout="wide")

st.markdown("""
<style>
    .header-box {
        background: linear-gradient(135deg, #1e1b4b 0%, #312e81 100%);
        color: white;
        padding: 22px;
        border-radius: 14px;
        margin-bottom: 20px;
    }
    .custom-table {
        width: 100%;
        border-collapse: collapse;
        margin-bottom: 25px;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    }
    .custom-table th.sec-header {
        background-color: #1f3845 !important;
        color: #ffffff !important;
        text-align: left;
        padding: 10px 14px;
        font-size: 15px;
        font-weight: bold;
        border: 1px solid #1f3845;
    }
    .custom-table th.col-header {
        background-color: #5a738e !important;
        color: #ffffff !important;
        padding: 8px 12px;
        font-size: 13px;
        font-weight: bold;
        border: 1px solid #4a637e;
    }
    .custom-table td {
        padding: 8px 12px;
        font-size: 13px;
        border: 1px solid #334155;
    }
    .custom-table tr.even-row td {
        background-color: #1e293b !important;
        color: #f8fafc;
    }
    .custom-table tr.odd-row td {
        background-color: #0f172a !important;
        color: #f8fafc;
    }
    .pos-val {
        color: #22c55e !important;
        font-weight: bold;
    }
    .neg-val {
        color: #ef4444 !important;
        font-weight: bold;
    }
    .total-row td {
        font-weight: bold;
        background-color: #1e293b !important;
        border-top: 2px solid #64748b;
        color: #ffffff;
    }
</style>
""", unsafe_allow_html=True)

# Helper function to convert any dirty string into a safe float
def clean_numeric_val(val):
    if val is None or pd.isna(val) or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    cleaned = re.sub(r'[^\d.-]', '', s)
    try:
        return float(cleaned) if cleaned else 0.0
    except Exception:
        return 0.0

def safe_currency(val):
    v = clean_numeric_val(val)
    return f"${v:,.2f}"

# -------------------------------------------------------------
# 1. Automatic PDF Parsing Engine (Pulls Sundays from POS PDF)
# -------------------------------------------------------------
@st.cache_data(show_spinner=False)
def parse_sundays_from_pdf_bytes(pdf_bytes):
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    records = []
    
    for page in reader.pages:
        text = page.extract_text()
        if not text:
            continue
            
        date_match = re.search(r'Date\s*[:\|]?\s*(\d{1,2}/\d{1,2}/\d{4})', text)
        if not date_match:
            date_match = re.search(r'(\d{2}/\d{2}/\d{4})', text)
            
        if date_match:
            date_str = date_match.group(1)
            try:
                dt = datetime.strptime(date_str, "%m/%d/%Y")
            except Exception:
                continue
            
            # Python weekday 6 = Sunday
            if dt.weekday() == 6:
                vape_val = 0.0
                vape_match = re.search(r'VAPE\s*[:\|]?\s*\$?\s*([0-9,]+\.\d{2})', text)
                if vape_match:
                    vape_val = float(vape_match.group(1).replace(',', ''))
                
                total_val = 0.0
                tender_split = text.split('Tender Collected')
                if len(tender_split) > 1:
                    above_tender = tender_split[0].strip()
                    lines = above_tender.splitlines()
                    last_line = lines[-1]
                    nums = re.findall(r'\$([0-9,]+\.\d{2})', last_line)
                    if nums:
                        total_val = float(nums[0].replace(',', ''))
                
                records.append({
                    'Date': dt.strftime('%m/%d/%Y'),
                    'Gross_Total': total_val,
                    'Vape_Sales': vape_val,
                    'Net_Alcohol': total_val - vape_val
                })
                
    df = pd.DataFrame(records).drop_duplicates(subset=['Date'])
    return df

# -------------------------------------------------------------
# 2. Resilient Google Sheet / CSV Loader
# -------------------------------------------------------------
SHEET_ID = "1NYSifzt4YDDZUkHviAD7bjBSFqbAJCR6vXBoXmHzIv0"
URL_CANDIDATES = [
    f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid=0",
    f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/gviz/tq?tqx=out:csv&gid=0",
    f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/gviz/tq?tqx=out:csv&sheet=Weekly%20KPIs"
]

def parse_kpi_text(raw_text):
    if "<!DOCTYPE html>" in raw_text or "<html" in raw_text.lower():
        return None, "Google returned an HTML login page instead of spreadsheet data."
    
    lines = raw_text.strip().splitlines()
    header_idx = None
    for idx, line in enumerate(lines[:20]):
        row = list(csv.reader([line]))[0] if line else []
        clean_tokens = [c.strip() for c in row]
        if 'Date' in clean_tokens:
            header_idx = idx
            break
            
    if header_idx is None:
        return None, "Could not locate 'Date' header row in Google Sheet data."
        
    df = pd.read_csv(io.StringIO(raw_text), skiprows=header_idx)
    df.columns = [str(c).strip() for c in df.columns]
    
    if 'Date' not in df.columns:
        return None, "Header found, but 'Date' column is missing."
        
    df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
    df = df.dropna(subset=['Date']).sort_values('Date').reset_index(drop=True)
    return df, None

@st.cache_data(ttl=300)
def fetch_online_kpis():
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/115.0.0.0 Safari/537.36"
    }
    for url in URL_CANDIDATES:
        try:
            resp = requests.get(url, headers=headers, timeout=12)
            if resp.status_code == 200 and "Date" in resp.text:
                return resp.text, None
        except Exception:
            continue
    return None, "Could not connect to Google Sheet via export endpoints."

st.sidebar.header("Data Settings")
csv_text, err_msg = fetch_online_kpis()
df_kpi = None

if csv_text:
    df_kpi, parse_err = parse_kpi_text(csv_text)
    if parse_err:
        st.sidebar.warning(parse_err)

manual_kpi = st.sidebar.file_uploader("Upload Weekly KPIs CSV (Backup)", type=["csv"])
if manual_kpi is not None:
    manual_text = manual_kpi.getvalue().decode('utf-8', errors='ignore')
    df_manual, manual_err = parse_kpi_text(manual_text)
    if df_manual is not None:
        df_kpi = df_manual
        st.sidebar.success("Loaded Weekly KPIs from uploaded file!")

if df_kpi is None:
    st.error("Cannot access Google Sheet. Please check link permissions or upload a backup CSV.")
    st.stop()

# Helper function to find and sum columns safely
def get_col_sum(df, possible_names):
    norm_targets = [re.sub(r'[^a-zA-Z0-9]', '', n).lower() for n in possible_names]
    for target in norm_targets:
        for col in df.columns:
            c_norm = re.sub(r'[^a-zA-Z0-9]', '', str(col)).lower()
            if c_norm == target:
                return float(df[col].apply(clean_numeric_val).sum())
    return 0.0

# -------------------------------------------------------------
# 3. Header Controls
# -------------------------------------------------------------
st.markdown("""
<div class='header-box'>
    <h2>LIQUORS VENTURES, LLC — Automated Executive Dashboard</h2>
    <p>Automated Quarterly Tracking & Apples-to-Apples Reconciliation from Weekly KPIs</p>
</div>
""", unsafe_allow_html=True)

c1, c2, c3 = st.columns([1, 2, 1])
with c1:
    year = st.selectbox("Year", [2026, 2025, 2024], index=0)
with c2:
    quarter = st.radio("Quarter", ["Q1", "Q2", "Q3", "Q4"], horizontal=True, index=2)
with c3:
    apples_to_apples = st.toggle("Apples-to-Apples ON", value=True)

# -------------------------------------------------------------
# 4. Permanent Archive Loader + Sidebar Fallback
# -------------------------------------------------------------
st.sidebar.header("Quarterly POS PDF")
uploaded_pdf = st.sidebar.file_uploader(
    f"Upload custom {year} {quarter} Daily POS PDF", 
    type=["pdf"]
)

df_sundays = pd.DataFrame()

possible_pdf_paths = [
    f"pdf_archive/{year}_{quarter}.pdf",
    f"sundays/{year}_{quarter}.pdf",
    f"{year}_{quarter}.pdf",
    f"{year}_{quarter.lower()}.pdf",
    f"pdf_archive/{year}_{quarter.lower()}.pdf"
]

found_archived_file = None
for p in possible_pdf_paths:
    if os.path.exists(p):
        found_archived_file = p
        break

if uploaded_pdf is not None:
    with st.spinner(f"Parsing uploaded {quarter} PDF..."):
        df_sundays = parse_sundays_from_pdf_bytes(uploaded_pdf.getvalue())
        st.sidebar.success(f"Loaded from upload: {len(df_sundays)} Sundays found!")
elif found_archived_file is not None:
    with open(found_archived_file, "rb") as f:
        file_bytes = f.read()
    df_sundays = parse_sundays_from_pdf_bytes(file_bytes)
    st.sidebar.info(f"Auto-loaded permanent archive: {found_archived_file} ({len(df_sundays)} Sundays)")

sunday_gross = df_sundays['Gross_Total'].sum() if not df_sundays.empty else 0.0
sunday_vape = df_sundays['Vape_Sales'].sum() if not df_sundays.empty else 0.0
sunday_net_deduction = df_sundays['Net_Alcohol'].sum() if not df_sundays.empty else 0.0

# -------------------------------------------------------------
# 5. Quarter Aggregation (13-Week Blocks)
# -------------------------------------------------------------
quarter_map = {"Q1": (1, 3), "Q2": (4, 6), "Q3": (7, 9), "Q4": (10, 12)}
start_m, end_m = quarter_map[quarter]

df_curr = df_kpi[(df_kpi['Date'].dt.year == year) & 
                 (df_kpi['Date'].dt.month >= start_m) & 
                 (df_kpi['Date'].dt.month <= end_m)].tail(13).copy()

df_prev = df_kpi[(df_kpi['Date'].dt.year == (year - 1)) & 
                 (df_kpi['Date'].dt.month >= start_m) & 
                 (df_kpi['Date'].dt.month <= end_m)].tail(13).copy()

# Pull inside sales w/o lottery
inside_names = ['Weekly Inside sales w/o lottery', 'INSIDE SALES']
gross_inside_curr = get_col_sum(df_curr, inside_names)
gross_inside_prev = get_col_sum(df_prev, inside_names)

shipping_curr = get_col_sum(df_curr, ['Total Shipping Sales'])
shipping_prev = get_col_sum(df_prev, ['Total Shipping Sales'])

# Pull Total Labor Under directly from sheet
total_labor_curr = get_col_sum(df_curr, ['Total Labor Under'])
if total_labor_curr == 0.0:
    total_labor_curr = get_col_sum(df_curr, ['Store Payroll 100hrs $2100']) + get_col_sum(df_curr, ['Shipping Payroll'])

total_labor_prev = get_col_sum(df_prev, ['Total Labor Under'])
if total_labor_prev == 0.0:
    total_labor_prev = get_col_sum(df_prev, ['Store Payroll 100hrs $2100']) + get_col_sum(df_prev, ['Shipping Payroll'])

vape_curr = get_col_sum(df_curr, ['VAPE SALES'])
vape_prev = get_col_sum(df_prev, ['VAPE SALES'])

total_sales_curr = gross_inside_curr + shipping_curr
total_sales_prev = gross_inside_prev + shipping_prev

trans_count_curr = get_col_sum(df_curr, ['INSIDE TRANS COUNT'])
parcels_curr = get_col_sum(df_curr, ['Total Parcels'])
if parcels_curr == 0:
    parcels_curr = len(df_curr) * 115

# Dynamic Normalization:
# Q1 & Q2 2026: Vape is deducted because vape launched late May 2025.
# Q3 2026 onwards: Vape is INCLUDED because both years had full vape operations.
if apples_to_apples and year == 2026:
    if quarter in ["Q1", "Q2"]:
        vape_deduction_curr = vape_curr
        vape_deduction_prev = vape_prev
        sunday_deduction_to_use = sunday_net_deduction
        vape_status_label = "Deducted (Vape Launched May '25)"
    else:
        vape_deduction_curr = 0.0
        vape_deduction_prev = 0.0
        sunday_deduction_to_use = sunday_gross
        vape_status_label = "Included (Active in Both Years)"

    adjusted_inside_curr = gross_inside_curr - vape_deduction_curr - sunday_deduction_to_use
    adjusted_inside_prev = gross_inside_prev - vape_deduction_prev
else:
    vape_deduction_curr = 0.0
    vape_deduction_prev = 0.0
    sunday_deduction_to_use = 0.0
    adjusted_inside_curr = gross_inside_curr
    adjusted_inside_prev = gross_inside_prev
    vape_status_label = "Unadjusted (Apples-to-Apples OFF)"

pct_inside_change = ((adjusted_inside_curr - adjusted_inside_prev) / adjusted_inside_prev) * 100 if adjusted_inside_prev != 0 else 0.0

# Total Labor Rate (Total Labor / Total Gross Sales)
labor_rate_curr = (total_labor_curr / total_sales_curr) * 100 if total_sales_curr != 0 else 0.0
labor_rate_prev = (total_labor_prev / total_sales_prev) * 100 if total_sales_prev != 0 else 0.0

# -------------------------------------------------------------
# 6. Bonus Calculations Engine
# -------------------------------------------------------------
def get_sales_bonus(growth):
    if growth >= 25.0:
        return 2500
    elif growth >= 20.0:
        return 2000
    elif growth >= 15.0:
        return 1500
    elif growth >= 8.0:
        return 1000
    elif growth >= 6.0:
        return 750
    elif growth >= 4.0:
        return 500
    return 0

def get_labor_bonus(rate):
    if rate <= 8.00:
        return 1000
    elif rate <= 9.00:
        return 750
    elif rate <= 11.00:
        return 500
    return 0

sales_bonus = get_sales_bonus(pct_inside_change)
labor_bonus = get_labor_bonus(labor_rate_curr)
total_bonus = sales_bonus + labor_bonus

# -------------------------------------------------------------
# 7. Render Tabs
# -------------------------------------------------------------
tab1, tab2, tab3 = st.tabs([
    "Executive Dashboard", 
    f"YoY Executive Tables ({quarter} {year} vs {quarter} {year-1})", 
    "Weekly Sales Data Breakdown"
])

# -------------------------------------------------------------
# TAB 1: VISUAL EXECUTIVE DASHBOARD
# -------------------------------------------------------------
with tab1:
    st.subheader("Apples-to-Apples Normalization Scorecard")
    s1, s2, s3, s4, s5 = st.columns(5)
    s1.metric("1. Gross Inside Sales", f"${gross_inside_curr:,.2f}", "Weekly Inside w/o Lottery")
    s2.metric("2. Vape Deductions", f"-${vape_deduction_curr:,.2f}", vape_status_label)
    s3.metric("3. Sunday Deductions", f"-${sunday_deduction_to_use:,.2f}", f"{len(df_sundays)} Sundays Detected")
    s4.metric("4. Adjusted Inside Sales", f"${adjusted_inside_curr:,.2f}", "Mon-Sat Core Store Sales")
    s5.metric("5. Normalized YoY", f"{pct_inside_change:+.2f}%", f"vs PY Baseline: ${adjusted_inside_prev:,.2f}")

    st.divider()

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("TOTAL GROSS SALES", f"${total_sales_curr:,.2f}", f"{((total_sales_curr - total_sales_prev)/total_sales_prev)*100:+.2f}% YoY" if total_sales_prev != 0 else "-")
    k2.metric("CORE ADJUSTED INSIDE", f"${adjusted_inside_curr:,.2f}", f"{pct_inside_change:+.2f}% Core")
    k3.metric("SHIPPING SALES (E-COMM)", f"${shipping_curr:,.2f}", f"{((shipping_curr - shipping_prev)/shipping_prev)*100:+.2f}% YoY" if shipping_prev else "-")
    k4.metric("TOTAL LABOR (STORE + SHIP)", f"${total_labor_curr:,.2f}", f"{((total_labor_curr - total_labor_prev)/total_labor_prev)*100:+.2f}% YoY" if total_labor_prev != 0 else "-")

    k5, k6, k7, k8 = st.columns(4)
    k5.metric("TOTAL LABOR % OF SALES", f"{labor_rate_curr:.2f}%", f"PY: {labor_rate_prev:.2f}%")
    k6.metric("STORE TRANSACTIONS", f"{int(trans_count_curr):,}", f"Avg Ticket: ${gross_inside_curr / max(1, trans_count_curr):.2f}")
    k7.metric("SUNDAY REVENUE EXTRACTED", f"${sunday_gross:,.2f}", f"Net Impact: -${sunday_net_deduction:,.2f}")
    k8.metric("TOTAL PARCELS SHIPPED", f"{int(parcels_curr):,}")

    st.divider()

    st.subheader("Quarterly Bonus Qualification")
    b1, b2, b3 = st.columns(3)
    b1.metric("Store Sales Bonus", f"${sales_bonus:,.2f}", f"Core Growth: {pct_inside_change:+.2f}%")
    b2.metric("Labor Rate Bonus", f"${labor_bonus:,.2f}", f"Total Labor Rate: {labor_rate_curr:.2f}%")
    b3.metric("Total Qualified Bonus", f"${total_bonus:,.2f}", "Combined Payout")

# -------------------------------------------------------------
# TAB 2: EXACT ALTERNATING COLOR TABLES + 1-SHEET EXCEL EXPORT
# -------------------------------------------------------------
with tab2:
    labor_diff = total_labor_curr - total_labor_prev
    labor_pct = (labor_diff / total_labor_prev) * 100 if total_labor_prev != 0 else 0.0

    t1_rows = [
        {
            "name": f"Total {quarter} Sales",
            "curr": total_sales_curr,
            "prev": total_sales_prev,
            "diff": total_sales_curr - total_sales_prev,
            "pct": ((total_sales_curr - total_sales_prev) / total_sales_prev) * 100 if total_sales_prev != 0 else None,
            "is_currency": True
        },
        {
            "name": "Net Inside Store Sales",
            "curr": adjusted_inside_curr,
            "prev": adjusted_inside_prev,
            "diff": adjusted_inside_curr - adjusted_inside_prev,
            "pct": pct_inside_change,
            "is_currency": True
        },
        {
            "name": "Shipping Revenue",
            "curr": shipping_curr,
            "prev": shipping_prev,
            "diff": shipping_curr - shipping_prev,
            "pct": ((shipping_curr - shipping_prev) / shipping_prev) * 100 if shipping_prev != 0 else None,
            "is_currency": True
        },
        {
            "name": f"{quarter} Total Labor (Store + Shipping Payroll)",
            "curr": total_labor_curr,
            "prev": total_labor_prev,
            "diff": labor_diff,
            "pct": labor_pct if total_labor_prev != 0 else None,
            "is_currency": True
        },
        {
            "name": "Labor Rate (Total Labor / Total Sales)",
            "curr": labor_rate_curr,
            "prev": labor_rate_prev,
            "diff": labor_rate_curr - labor_rate_prev,
            "pct": None,
            "is_currency": False,
            "is_rate": True
        },
        {
            "name": "Total Qualified Bonus",
            "curr": total_bonus,
            "prev": None,
            "diff": None,
            "pct": None,
            "is_currency": True,
            "custom_note": f"Sales: ${sales_bonus:,.2f} - Labor:${labor_bonus:,.2f}"
        }
    ]

    vape_diff = vape_deduction_curr - vape_deduction_prev
    vape_pct = (vape_diff / vape_deduction_prev) * 100 if vape_deduction_prev != 0 else None

    t2_rows = [
        {
            "name": "Inside Sales (Gross)",
            "curr": gross_inside_curr,
            "prev": gross_inside_prev,
            "diff": gross_inside_curr - gross_inside_prev,
            "pct": ((gross_inside_curr - gross_inside_prev) / gross_inside_prev) * 100 if gross_inside_prev != 0 else None,
            "is_currency": True
        },
        {
            "name": "Vape Sales Deductions",
            "curr": -vape_deduction_curr,
            "prev": -vape_deduction_prev,
            "diff": -vape_diff,
            "pct": vape_pct,
            "is_currency": True
        },
        {
            "name": "Sunday Sales Deductions",
            "curr": -sunday_deduction_to_use,
            "prev": 0.0,
            "diff": -sunday_deduction_to_use,
            "pct": None,
            "is_currency": True
        },
        {
            "name": "Total Net Inside Store Sales",
            "curr": adjusted_inside_curr,
            "prev": adjusted_inside_prev,
            "diff": adjusted_inside_curr - adjusted_inside_prev,
            "pct": pct_inside_change,
            "is_currency": True
        }
    ]

    def render_clean_html_table(title_text, headers, rows_data):
        lines = []
        lines.append('<table class="custom-table">')
        lines.append('<thead>')
        lines.append('<tr>')
        lines.append(f'<th colspan="{len(headers)}" class="sec-header">{title_text}</th>')
        lines.append('</tr>')
        lines.append('<tr>')
        for h in headers:
            align = "left" if h in ["Metric", "Category", "Date"] else "right"
            lines.append(f'<th class="col-header" style="text-align:{align};">{h}</th>')
        lines.append('</tr>')
        lines.append('</thead>')
        lines.append('<tbody>')
        
        for idx, row in enumerate(rows_data):
            row_cls = "even-row" if idx % 2 == 0 else "odd-row"
            lines.append(f'<tr class="{row_cls}">')
            lines.append(f'<td style="text-align:left; font-weight:bold;">{row["name"]}</td>')
            
            if row.get("is_rate"):
                curr_str = f"{row['curr']:.2f}%"
            elif row.get("is_currency"):
                curr_str = f"\${row['curr']:,.2f}" if row['curr'] is not None else "-"
            else:
                curr_str = str(row['curr'])
            lines.append(f'<td style="text-align:right;">{curr_str}</td>')
            
            if row.get("is_rate"):
                prev_str = f"{row['prev']:.2f}%"
            elif row.get("is_currency"):
                prev_str = f"\${row['prev']:,.2f}" if row['prev'] is not None else "-"
            else:
                prev_str = str(row['prev']) if row['prev'] is not None else "-"
            lines.append(f'<td style="text-align:right;">{prev_str}</td>')
            
            if row.get("custom_note"):
                lines.append(f'<td style="text-align:right;">{row["custom_note"]}</td>')
            elif row.get("is_rate"):
                val = row['diff']
                cls = "pos-val" if val > 0 else ("neg-val" if val < 0 else "")
                lines.append(f'<td style="text-align:right;" class="{cls}">{val:+.2f}% pts</td>')
            elif row['diff'] is not None:
                val = row['diff']
                cls = "pos-val" if val > 0 else ("neg-val" if val < 0 else "")
                lines.append(f'<td style="text-align:right;" class="{cls}">\${val:+,.2f}</td>')
            else:
                lines.append('<td style="text-align:right;">-</td>')
                
            if row['pct'] is not None:
                val = row['pct']
                cls = "pos-val" if val > 0 else ("neg-val" if val < 0 else "")
                lines.append(f'<td style="text-align:right;" class="{cls}">{val:+.2f}%</td>')
            else:
                lines.append('<td style="text-align:right;">-</td>')
                
            lines.append('</tr>')
        lines.append('</tbody>')
        lines.append('</table>')
        return "".join(lines)

    headers_5col = ["Metric", f"{quarter} {year}", f"{quarter} {year-1}", "Gain/Loss (\$)", "Gain/Loss (%)"]
    st.html(render_clean_html_table(f"1. EXECUTIVE OVERVIEW ({quarter} {year} vs {quarter} {year-1})", headers_5col, t1_rows))

    headers_t2 = ["Category", f"{quarter} {year}", f"{quarter} {year-1}", "Gain/Loss (\$)", "Gain/Loss (%)"]
    st.html(render_clean_html_table("2. INSIDE SALES PERFORMANCE (YOY)", headers_t2, t2_rows))

    # 3. Sunday Sales Deep Dive Table
    if not df_sundays.empty:
        sun_headers = ["Date", "Gross Sunday Total", "Vape Portion", "Net Sunday Impact"]
        slines = []
        slines.append('<table class="custom-table">')
        slines.append('<thead>')
        slines.append('<tr>')
        slines.append(f'<th colspan="4" class="sec-header">3. {quarter} {year} SUNDAY SALES DEEP DIVE</th>')
        slines.append('</tr>')
        slines.append('<tr>')
        for h in sun_headers:
            align = "left" if h == "Date" else "right"
            slines.append(f'<th class="col-header" style="text-align:{align};">{h}</th>')
        slines.append('</tr>')
        slines.append('</thead>')
        slines.append('<tbody>')
        
        for idx, r in df_sundays.iterrows():
            row_cls = "even-row" if idx % 2 == 0 else "odd-row"
            slines.append(f'<tr class="{row_cls}">')
            slines.append(f'<td style="text-align:left;">{r["Date"]}</td>')
            slines.append(f'<td style="text-align:right;">\${r["Gross_Total"]:,.2f}</td>')
            slines.append(f'<td style="text-align:right;">\${r["Vape_Sales"]:,.2f}</td>')
            slines.append(f'<td style="text-align:right;">\${r["Net_Alcohol"]:,.2f}</td>')
            slines.append('</tr>')
            
        slines.append('<tr class="total-row">')
        slines.append('<td style="text-align:left;">TOTAL</td>')
        slines.append(f'<td style="text-align:right;">\${sunday_gross:,.2f}</td>')
        slines.append(f'<td style="text-align:right;">\${sunday_vape:,.2f}</td>')
        slines.append(f'<td style="text-align:right;">\${sunday_net_deduction:,.2f}</td>')
        slines.append('</tr>')
        slines.append('</tbody>')
        slines.append('</table>')
        
        st.html("".join(slines))
    else:
        st.info(f"Drop `{year}_{quarter}.pdf` into your repository or the sidebar to populate the 13-week Sunday table!")

    # -------------------------------------------------------------
    # BUILD SINGLE-SHEET EXCEL (.XLSX) REPLICATING ATTACHMENT 2
    # -------------------------------------------------------------
    st.divider()

    def build_executive_excel():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{quarter} {year} Executive Report"
        ws.views.sheetView[0].showGridLines = True

        title_fill = PatternFill(start_color="1E3A47", end_color="1E3A47", fill_type="solid")
        title_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        header_fill = PatternFill(start_color="5A738E", end_color="5A738E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        row_even_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
        row_odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        font_bold = Font(name="Calibri", size=11, bold=True)
        font_regular = Font(name="Calibri", size=11)
        font_green = Font(name="Calibri", size=11, bold=True, color="008000")
        font_red = Font(name="Calibri", size=11, bold=True, color="CC0000")

        thin_border = Border(
            left=Side(style='thin', color='D0D7DE'),
            right=Side(style='thin', color='D0D7DE'),
            top=Side(style='thin', color='D0D7DE'),
            bottom=Side(style='thin', color='D0D7DE')
        )

        align_left = Alignment(horizontal='left', vertical='center')
        align_right = Alignment(horizontal='right', vertical='center')

        cur_row = 1

        def write_section(title, headers, rows):
            nonlocal cur_row
            ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=len(headers))
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=cur_row, column=c)
                cell.fill = title_fill
                cell.font = title_font
                cell.alignment = align_left
                cell.border = thin_border
            ws.cell(row=cur_row, column=1, value=title)
            cur_row += 1

            for c_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=cur_row, column=c_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_left if c_idx == 1 else align_right
                cell.border = thin_border
            cur_row += 1

            for r_idx, r in enumerate(rows):
                fill = row_even_fill if r_idx % 2 == 0 else row_odd_fill
                
                c1 = ws.cell(row=cur_row, column=1, value=r["name"])
                c1.fill = fill
                c1.font = font_bold
                c1.alignment = align_left
                c1.border = thin_border

                c2 = ws.cell(row=cur_row, column=2, value=r["curr"])
                c2.fill = fill
                c2.font = font_regular
                c2.alignment = align_right
                c2.border = thin_border
                if r.get("is_rate"):
                    c2.number_format = '0.00%'
                    c2.value = r["curr"] / 100.0 if r["curr"] is not None else ""
                elif r.get("is_currency"):
                    c2.number_format = '\$#,##0.00'

                c3 = ws.cell(row=cur_row, column=3, value=r["prev"])
                c3.fill = fill
                c3.font = font_regular
                c3.alignment = align_right
                c3.border = thin_border
                if r.get("is_rate"):
                    c3.number_format = '0.00%'
                    c3.value = r["prev"] / 100.0 if r["prev"] is not None else ""
                elif r.get("is_currency"):
                    c3.number_format = '\$#,##0.00'

                c4 = ws.cell(row=cur_row, column=4)
                c4.fill = fill
                c4.alignment = align_right
                c4.border = thin_border
                if r.get("custom_note"):
                    c4.value = r["custom_note"]
                    c4.font = font_regular
                elif r.get("is_rate"):
                    c4.value = f"{r['diff']:+.2f}% pts"
                    c4.font = font_green if r['diff'] > 0 else (font_red if r['diff'] < 0 else font_regular)
                elif r['diff'] is not None:
                    c4.value = r['diff']
                    c4.number_format = '$#,##0.00;[Red]($#,##0.00);"-"'
                    c4.font = font_green if r['diff'] > 0 else (font_red if r['diff'] < 0 else font_regular)
                else:
                    c4.value = "-"
                    c4.font = font_regular

                c5 = ws.cell(row=cur_row, column=5)
                c5.fill = fill
                c5.alignment = align_right
                c5.border = thin_border
                if r['pct'] is not None:
                    c5.value = r['pct'] / 100.0
                    c5.number_format = '+0.00%;-0.00%;0.00%'
                    c5.font = font_green if r['pct'] > 0 else (font_red if r['pct'] < 0 else font_regular)
                else:
                    c5.value = "-"
                    c5.font = font_regular

                cur_row += 1

            cur_row += 2

        write_section(f"1. EXECUTIVE OVERVIEW ({quarter} {year} vs {quarter} {year-1})", headers_5col, t1_rows)
        write_section("2. INSIDE SALES PERFORMANCE (YOY)", headers_t2, t2_rows)

        if not df_sundays.empty:
            sun_headers = ["Date", "Gross Sunday Total", "Vape Portion", "Net Sunday Impact"]
            ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=4)
            for c in range(1, 5):
                cell = ws.cell(row=cur_row, column=c)
                cell.fill = title_fill
                cell.font = title_font
                cell.alignment = align_left
                cell.border = thin_border
            ws.cell(row=cur_row, column=1, value=f"3. {quarter} {year} SUNDAY SALES DEEP DIVE")
            cur_row += 1

            for c_idx, h in enumerate(sun_headers, 1):
                cell = ws.cell(row=cur_row, column=c_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_left if c_idx == 1 else align_right
                cell.border = thin_border
            cur_row += 1

            for s_idx, sr in df_sundays.iterrows():
                fill = row_even_fill if s_idx % 2 == 0 else row_odd_fill
                c1 = ws.cell(row=cur_row, column=1, value=sr['Date'])
                c1.fill = fill
                c1.font = font_regular
                c1.alignment = align_left
                c1.border = thin_border

                for ci, k in enumerate(['Gross_Total', 'Vape_Sales', 'Net_Alcohol'], 2):
                    cell = ws.cell(row=cur_row, column=ci, value=sr[k])
                    cell.fill = fill
                    cell.font = font_regular
                    cell.alignment = align_right
                    cell.border = thin_border
                    cell.number_format = '\$#,##0.00'
                cur_row += 1

            tot_row = ws.cell(row=cur_row, column=1, value="TOTAL")
            tot_row.font = font_bold
            tot_row.alignment = align_left
            tot_row.fill = row_even_fill
            tot_row.border = thin_border

            for ci, val in enumerate([sunday_gross, sunday_vape, sunday_net_deduction], 2):
                cell = ws.cell(row=cur_row, column=ci, value=val)
                cell.font = font_bold
                cell.fill = row_even_fill
                cell.alignment = align_right
                cell.border = thin_border
                cell.number_format = '\$#,##0.00'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    excel_bytes = build_executive_excel()

    st.download_button(
        label=f"📥 Download {quarter} {year} All Executive Tables (.xlsx)",
        data=excel_bytes,
        file_name=f"Liquors_Ventures_{quarter}_{year}_Executive_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# -------------------------------------------------------------
# TAB 3: WEEKLY SALES DATA BREAKDOWN WITH TOTALS
# -------------------------------------------------------------
with tab3:
    st.markdown(f"### 13-Week Data Points from Weekly KPIs ({year} vs {year-1})")
    
    col_a, col_b = st.columns(2)
    numeric_display = [
        inside_names[0], 'Total Shipping Sales', 'Total Weekly Sales with out Lottery', 
        'VAPE SALES', 'Store Payroll 100hrs \$2100', 'Shipping Payroll', 'Total Labor Under'
    ]
    
    with col_a:
        st.markdown(f"#### {quarter} {year} Weekly Line Items")
        show_cols_curr = [c for c in ['Date'] + numeric_display if c in df_curr.columns]
        df_c = df_curr[show_cols_curr].copy()
        df_c['Date'] = df_c['Date'].dt.strftime('%Y-%m-%d')
        
        total_dict_curr = {'Date': 'TOTAL'}
        for col in numeric_display:
            if col in df_curr.columns:
                total_dict_curr[col] = float(df_curr[col].apply(clean_numeric_val).sum())
        df_c = pd.concat([df_c, pd.DataFrame([total_dict_curr])], ignore_index=True)
        
        for col in numeric_display:
            if col in df_c.columns:
                df_c[col] = df_c[col].apply(safe_currency)
                
        st.dataframe(df_c, width="stretch", hide_index=True)
        
    with col_b:
        st.markdown(f"#### {quarter} {year-1} Weekly Line Items")
        show_cols_prev = [c for c in ['Date'] + numeric_display if c in df_prev.columns]
        df_p = df_prev[show_cols_prev].copy()
        df_p['Date'] = df_p['Date'].dt.strftime('%Y-%m-%d')
        
        total_dict_prev = {'Date': 'TOTAL'}
        for col in numeric_display:
            if col in df_prev.columns:
                total_dict_prev[col] = float(df_prev[col].apply(clean_numeric_val).sum())
        df_p = pd.concat([df_p, pd.DataFrame([total_dict_prev])], ignore_index=True)
        
        for col in numeric_display:
            if col in df_p.columns:
                df_p[col] = df_p[col].apply(safe_currency)
                
        st.dataframe(df_p, width="stretch", hide_index=True)
